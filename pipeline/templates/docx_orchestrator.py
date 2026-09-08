#!/usr/bin/env python3
"""Оркестратор правок .docx через python-docx (канон §4 docx-protocol.md)."""
import re
import shutil
from docx import Document
from docx.shared import Pt, Mm, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FILE = "результат.docx"


def set_run_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    """Установить шрифт run (включая cs и eastAsia)."""
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    for attr in ['w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia']:
        rFonts.set(qn(attr), name)


def set_para_format(para, alignment=None, first_line_indent_cm=None,
                    line_spacing=None, bold=None, font_name="Times New Roman", font_size=12):
    """Установить формат параграфа и всех его runs."""
    if alignment is not None:
        para.alignment = alignment
    if first_line_indent_cm is not None:
        para.paragraph_format.first_line_indent = Cm(first_line_indent_cm)
    if line_spacing is not None:
        para.paragraph_format.line_spacing = line_spacing
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.space_before = Pt(0)
    for run in para.runs:
        set_run_font(run, font_name, font_size, bold=bold)


def set_cell_format(cell, font_name="Times New Roman", font_size=12):
    """Установить шрифт во всех параграфах ячейки таблицы."""
    for para in cell.paragraphs:
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.space_before = Pt(0)
        for run in para.runs:
            set_run_font(run, font_name, font_size)


def find_para(doc, text_startswith=None, text_contains=None):
    """Найти параграф по началу текста или содержимому."""
    for i, p in enumerate(doc.paragraphs):
        if text_startswith and p.text.startswith(text_startswith):
            return i, p
        if text_contains and text_contains in p.text:
            return i, p
    return None, None


def insert_heading_before(doc, anchor_text, heading_text, check_duplicate=True):
    """Вставить заголовок перед абзацем-якорем (см. §1.6 docx-protocol.md)."""
    if check_duplicate:
        for i, p in enumerate(doc.paragraphs):
            if heading_text.strip() in p.text:
                print(f"WARN: заголовок '{heading_text}' уже существует в P{i} — пропускаем вставку")
                return False
    _, anchor = find_para(doc, text_startswith=anchor_text)
    if anchor is None:
        print(f"WARN: якорь '{anchor_text}' не найден")
        return False
    new_p = OxmlElement('w:p')
    new_pPr = OxmlElement('w:pPr')
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    new_pPr.append(jc)
    new_p.append(new_pPr)
    new_r = OxmlElement('w:r')
    new_rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    for attr in ['w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia']:
        rFonts.set(qn(attr), 'Times New Roman')
    new_rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '24')
    new_rPr.append(sz)
    b = OxmlElement('w:b')
    new_rPr.append(b)
    new_r.append(new_rPr)
    t = OxmlElement('w:t')
    t.text = heading_text
    new_r.append(t)
    new_p.append(new_r)
    anchor._element.addprevious(new_p)
    return True


def replace_cross_run_text(paragraph, old, new):
    """Заменить подстроку в параграфе с учётом cross-run разбивки (см. §5.1)."""
    if not paragraph.runs:
        return False
    full_text = paragraph.text
    if old not in full_text:
        return False
    new_text = full_text.replace(old, new)
    for run in paragraph.runs[1:]:
        run.text = ""
    paragraph.runs[0].text = new_text
    return True


def replace_cross_run_in_cell(cell, old, new):
    """Заменить подстроку во всех параграфах ячейки таблицы (cross-run aware)."""
    found = False
    for para in cell.paragraphs:
        if replace_cross_run_text(para, old, new):
            found = True
    return found


def renumber_references(doc, ref_start_pattern='ЛИТЕРАТУРА', ref_end_pattern='СВЕДЕНИЯ ОБ АВТОРАХ'):
    """Проставить нумерацию 1..N (см. §1.8: regex, не срезы строк)."""
    ref_start = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == ref_start_pattern:
            ref_start = i + 1
            break
    if ref_start is None:
        print(f"WARN: заголовок '{ref_start_pattern}' не найден")
        return 0
    n = 1
    for i in range(ref_start, len(doc.paragraphs)):
        p = doc.paragraphs[i]
        txt = p.text.strip()
        if not txt:
            continue
        if txt == ref_end_pattern:
            break
        cleaned = re.sub(r'^\s*\d+\.\s*', '', txt)
        new_text = f"{n}. {cleaned}"
        for run in p.runs[1:]:
            run.text = ""
        if p.runs:
            p.runs[0].text = new_text
        else:
            p.text = new_text
        n += 1
    return n - 1


if __name__ == "__main__":
    import sys
    src = sys.argv[1] if len(sys.argv) > 1 else "исходник.docx"
    out = sys.argv[2] if len(sys.argv) > 2 else FILE
    shutil.copy(src, out)
    doc = Document(out)
    for section in doc.sections:
        section.top_margin = Mm(25)
        section.bottom_margin = Mm(25)
        section.left_margin = Mm(25)
        section.right_margin = Mm(25)
    for para in doc.paragraphs:
        set_para_format(para, line_spacing=1.0, first_line_indent_cm=1.25, font_size=12)
    doc.save(out)
    print(f"Сохранено: {out}")
